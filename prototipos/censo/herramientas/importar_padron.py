"""Convierte el Excel del Libro de Padrón (transcripción de las fichas) en los
documentos de la base: un asociado por ficha y un stand por cada N° de stand.

    pip install openpyxl
    python3 herramientas/importar_padron.py Padron_Asociados.xlsx SALIDA/

Escribe SALIDA/asociados/<id>.json, SALIDA/stands/<codigo>.json, SALIDA/semilla.json
(todo junto, para probar en local) y SALIDA/resumen.json (avisos para revisar).
Los datos personales se quedan fuera del repositorio: SALIDA no se versiona.
"""
import sys, os, re, json
import openpyxl

XLSX, SALIDA = sys.argv[1], sys.argv[2]
AQUI = os.path.dirname(os.path.abspath(__file__))
FECHA_CARGA = "2026-10-01T09:00:00-05:00"

# Galería de cada stand según el plano real (src/plano.js).
plano_js = open(os.path.join(AQUI, "..", "src", "plano.js"), encoding="utf-8").read()
PLANO = json.loads(plano_js[plano_js.index("{"): plano_js.rindex("}") + 1])
GALERIA = {c: g for c, g, *_ in PLANO["stands"]}

MINUSCULAS = {"de", "del", "la", "las", "los", "y", "e", "da", "van", "von"}
def titulo(s):
    """'ALEJANDRINA ACHULLI MAIHURE' → 'Alejandrina Achulli Maihure' (respeta 'de', 'del'…)."""
    s = re.sub(r"\s+", " ", str(s or "").strip())
    palabras = []
    for i, p in enumerate(s.lower().split(" ")):
        palabras.append(p if i and p in MINUSCULAS else p[:1].upper() + p[1:])
    return " ".join(palabras)

def texto(v):
    return re.sub(r"\s+", " ", str(v).strip()) if v is not None else ""

def fecha_iso(v):
    """'10/03/1966', '24-11-1991' o una fecha de Excel (datetime) → '1966-03-10'. Si no se puede leer, ''."""
    if hasattr(v, "year") and hasattr(v, "month"):  # celda con formato de fecha
        return f"{v.year:04d}-{v.month:02d}-{v.day:02d}" if 1900 <= v.year <= 2100 else ""
    m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", texto(v))
    if not m:
        return ""
    d, mes, a = map(int, m.groups())
    if not (1900 <= a <= 2100 and 1 <= mes <= 12 and 1 <= d <= 31):
        return ""
    return f"{a:04d}-{mes:02d}-{d:02d}"

ESTADO_CIVIL = {"SOLTERO(A)": "Soltero(a)", "SOLTERA": "Soltero(a)", "SOLTERO": "Soltero(a)", "CASADO(A)": "Casado(a)",
                "VIUDO(A)": "Viudo(a)", "DIVORCIADO(A)": "Divorciado(a)", "CONVIVIENTE": "Conviviente"}
INSTRUCCION = {"PRIMARIA": "Primaria", "SECUNDARIA": "Secundaria", "SUPERIOR": "Superior", "TECNICO": "Técnico",
               "UNIVERSITARIA": "Universitaria", "MAGISTER": "Maestría", "ILETRADA": "Sin instrucción"}
DISTRITOS = {"SJM": "San Juan de Miraflores", "SJL": "San Juan de Lurigancho", "SMP": "San Martín de Porres",
             "VES": "Villa El Salvador", "VMT": "Villa María del Triunfo", "SMP.": "San Martín de Porres",
             "STA. ANITA": "Santa Anita", "STA ANITA": "Santa Anita", "STGO. SURCO": "Santiago de Surco",
             "RIMAC": "Rímac", "EEUU": "EE. UU.", "LIMA": "Cercado de Lima"}

def distrito(v):
    t = texto(v)
    return DISTRITOS.get(t.upper(), titulo(t))

wb = openpyxl.load_workbook(XLSX, data_only=True)
hoja = wb.worksheets[0]
cab = [texto(c) for c in next(hoja.iter_rows(min_row=4, max_row=4, values_only=True))]
col = {nombre: i for i, nombre in enumerate(cab)}
filas = [r for r in hoja.iter_rows(min_row=5, values_only=True) if r[0] is not None]


def normal(s):
    """Nombre normalizado para comparar (sin tildes, mayúsculas, espacios simples)."""
    import unicodedata
    t = unicodedata.normalize("NFD", texto(s)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", t).strip()

PARENTESCO = {"MAMA": "Madre", "MADRE": "Madre", "PAPA": "Padre", "PADRE": "Padre", "HRNO": "Hermano(a)", "HNO": "Hermano(a)", "HERMANO": "Hermano(a)",
              "HRNA": "Hermano(a)", "HNA": "Hermano(a)", "HERMANA": "Hermano(a)", "SOBRINO": "Sobrino(a)", "SOBRINA": "Sobrino(a)", "NIETO": "Nieto(a)", "NIETA": "Nieto(a)"}
ESTUDIOS = {"PRIMARIA": "Primaria", "PRI": "Primaria", "PRIM": "Primaria", "SECUNDARIA": "Secundaria", "SEC": "Secundaria", "SECUND": "Secundaria",
            "INICIAL": "Inicial", "INI": "Inicial", "NINGUNO": "Ninguno", "SIN ESTUDIOS": "Ninguno", "SUP": "Superior técnica", "SUP TEC": "Superior técnica",
            "UNIV": "Universitaria", "UNI": "Universitaria",
            "TECNICO": "Superior técnica", "TECNICA": "Superior técnica", "SUPERIOR": "Superior técnica", "SUPERIOR TECNICA": "Superior técnica",
            "UNIVERSITARIA": "Universitaria", "UNIVERSITARIO": "Universitaria", "UNIVERSIDAD": "Universitaria"}

def edad_txt(v):
    m = re.search(r"\d{1,3}", texto(v))
    return m.group(0) if m else ""

def estudios_txt(v):
    t = normal(v)
    return ESTUDIOS.get(t, titulo(t)) if t else ""

asociados, stands, avisos = {}, {}, []
sin_numero = 0
for r in filas:
    v = lambda nombre: r[col[nombre]]
    numero = v("N° Asociado")
    if isinstance(numero, int):
        numero_txt, ident = str(numero), f"a{numero:03d}"
    else:
        sin_numero += 1
        numero_txt, ident = f"S/N-{sin_numero}", f"asn{sin_numero}"
    fono = re.sub(r"\D", "", texto(v("Celular / Fijo")))
    celular = fono if re.fullmatch(r"9\d{8}", fono) else ""
    telefono = "" if celular else fono
    codigos = re.findall(r"\d{4}", texto(v("N° Stand(s)")))
    cuentas = [c.strip() for c in re.split(r"\s*-\s*", texto(v("N° Cuenta Bco."))) if c.strip()]
    nota = texto(v("Observaciones"))
    nac, ing = fecha_iso(v("Fecha Nac.")), fecha_iso(v("Fecha de Ingreso"))
    if texto(v("Fecha de Ingreso")) and not ing:
        avisos.append(f"N° {numero_txt}: fecha de ingreso ilegible ({texto(v('Fecha de Ingreso'))}).")
    conyuge = titulo(v("Cónyuge / Conviviente"))
    asociados[ident] = {
        "numero": numero_txt,
        "nombres": titulo(v("Nombres")),
        "apellidoPaterno": titulo(v("Apellido Paterno")),
        "apellidoMaterno": titulo(str(v("Apellido Materno") or "").rstrip(",")),
        "dni": re.sub(r"\D", "", texto(v("DNI"))),
        "fechaNacimiento": nac,
        "nacimiento": {"departamento": titulo(v("Departamento")), "provincia": titulo(v("Provincia")), "distrito": titulo(v("Distrito (nac.)"))},
        "ocupacion": titulo(v("Ocupación")),
        "instruccion": INSTRUCCION.get(texto(v("Grado de Instrucción")).upper(), titulo(v("Grado de Instrucción"))),
        "estadoCivil": ESTADO_CIVIL.get(texto(v("Estado Civil")).upper(), ""),
        "direccion": texto(v("Dirección")),
        "distritoResidencia": distrito(v("Distrito (dirección)")),
        "celular": celular,
        "telefono": telefono,
        "correo": texto(v("Correo")).replace(" ", "").lower(),
        "fechaIngreso": ing,
        "estado": "activo",
        "cuentaBancaria": texto(v("N° Cuenta Bco.")),
        "conyuge": {"nombre": conyuge, "dni": "", "celular": ""} if conyuge else None,
        "hijos": [], "familiares": [],
        "censo": {"estado": "pendiente", "fecha": None, "visita": None},
        "observaciones": f"Revisar (transcripción del libro): {nota}" if nota else "",
        "revisar": bool(nota),
        "archivos": {}, "documentos": [],
        "compromiso": {"estatutos": False, "datos": False},
        "origen": {"libro": "Libro de Padrón de Asociados", "lote": v("Lote"), "foto": texto(v("Foto original")), "carpeta": texto(v("Carpeta"))},
        "creadoAt": FECHA_CARGA, "actualizadoAt": FECHA_CARGA,
    }
    nota_ficha = texto(v("Observaciones de la ficha")) if "Observaciones de la ficha" in col else ""
    if nota_ficha:
        asociados[ident]["observaciones"] = " ".join(x for x in [asociados[ident]["observaciones"], f"Nota de la ficha: {nota_ficha}"] if x)
        asociados[ident]["revisar"] = True
    for i, c in enumerate(codigos):
        if c not in GALERIA:
            avisos.append(f"N° {numero_txt}: el stand {c} no está en el plano.")
            continue
        if c in stands:
            # Dos fichas con el mismo stand (posible traspaso): queda la que tiene N° de asociado.
            previo = stands[c]["propietarioId"]
            gana = previo if previo.startswith("a") and not previo.startswith("asn") else ident
            pierde = ident if gana == previo else previo
            stands[c]["propietarioId"] = gana
            avisos.append(f"Stand {c}: figura en dos fichas ({asociados[gana]['numero']} y {asociados[pierde]['numero']}); "
                          f"quedó a nombre de N° {asociados[gana]['numero']}. Si se vendió, regístralo con «Registrar venta».")
            nota_extra = f"Figura en la ficha del stand {c}, que hoy está a nombre de N° {asociados[gana]['numero']}: verificar traspaso."
            asociados[pierde]["observaciones"] = " ".join(x for x in [asociados[pierde]["observaciones"], nota_extra] if x)
            asociados[pierde]["revisar"] = True
            continue
        stands[c] = {
            "codigo": c, "galeria": GALERIA[c], "area": None, "giro": "",
            "propietarioId": ident, "inquilino": None, "estado": "propietario",
            "cuenta": cuentas[i] if i < len(cuentas) else "",
        }


# ---- Hojas "Hijos" y "Otros familiares" (por N° de asociado; los S/N se emparejan por nombre) ----
por_numero = {a["numero"]: k for k, a in asociados.items()}
por_nombre = {normal(f"{a['nombres']} {a['apellidoPaterno']} {a['apellidoMaterno']}"): k for k, a in asociados.items()}

def ident_de(numero_v, nombre_v, hoja_nombre):
    n = texto(numero_v)
    if n and n.upper() != "S/N" and n in por_numero:
        return por_numero[n]
    k = por_nombre.get(normal(nombre_v))
    if k is None:
        avisos.append(f"Hoja {hoja_nombre}: no se encontró al asociado N° {n} «{texto(nombre_v)}».")
    return k

def filas_hoja(nombre):
    if nombre not in wb.sheetnames:
        return []
    h = wb[nombre]
    cab = [texto(c) for c in next(h.iter_rows(min_row=1, max_row=1, values_only=True))]
    return [dict(zip(cab, r)) for r in h.iter_rows(min_row=2, values_only=True) if r and r[0] not in (None, "")]

n_hijos = n_fam = 0
for r in filas_hoja("Hijos"):
    k = ident_de(r.get("N° Asociado"), r.get("Asociado"), "Hijos")
    if not k or not texto(r.get("Nombre del hijo/a")):
        continue
    asociados[k]["hijos"].append({"nombre": titulo(r.get("Nombre del hijo/a")), "edad": edad_txt(r.get("Edad")),
                                  "estudios": estudios_txt(r.get("Estudios")), "dni": re.sub(r"\D", "", texto(r.get("DNI hijo")))})
    n_hijos += 1
for r in filas_hoja("Otros familiares"):
    k = ident_de(r.get("N° Asociado"), r.get("Asociado"), "Otros familiares")
    if not k or not texto(r.get("Nombre")):
        continue
    par = normal(r.get("Parentesco"))
    asociados[k]["familiares"].append({"nombre": titulo(r.get("Nombre")), "parentesco": PARENTESCO.get(par, titulo(par)) if par else "",
                                       "edad": edad_txt(r.get("Edad")), "estudios": estudios_txt(r.get("Estudios"))})
    n_fam += 1
print(f"Hijos leídos: {n_hijos} · Otros familiares: {n_fam}")

os.makedirs(os.path.join(SALIDA, "asociados"), exist_ok=True)
os.makedirs(os.path.join(SALIDA, "stands"), exist_ok=True)
semilla = []
for ident, a in asociados.items():
    json.dump(a, open(os.path.join(SALIDA, "asociados", f"{ident}.json"), "w"), ensure_ascii=False)
    semilla.append({"path": f"asociados/{ident}", "data": a})
for c, s in stands.items():
    json.dump(s, open(os.path.join(SALIDA, "stands", f"{c}.json"), "w"), ensure_ascii=False)
    semilla.append({"path": f"stands/{c}", "data": s})
json.dump(semilla, open(os.path.join(SALIDA, "semilla.json"), "w"), ensure_ascii=False)
resumen = {"asociados": len(asociados), "stands": len(stands), "por_revisar": sum(a["revisar"] for a in asociados.values()),
           "sin_stand_en_plano": len(GALERIA) - len(stands), "avisos": avisos}
json.dump(resumen, open(os.path.join(SALIDA, "resumen.json"), "w"), ensure_ascii=False, indent=1)
print(json.dumps(resumen, ensure_ascii=False, indent=1))
