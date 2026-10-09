#!/usr/bin/env python3
"""Genera el dashboard HTML de jubilaciones a partir del
"Reporte_Control_MaestraEstatus" (padrón de asociados activos AMSP).

Uso:
    python3 build.py <Reporte_Control_MaestraEstatusYYYYMMDD.xlsx> [salida.html]

Lee la hoja `DetMaeSituAct` (cabecera en la fila 5), compacta los campos que
usa el dashboard (código, nombre, DNI, sexo, FECNAC, edad reportada, dpto.,
unidad ejecutora, escala, estado, clasificación AMSP) y los incrusta en
`plantilla.html`. El cálculo de edades y jubilaciones se hace en el navegador
con la fecha de corte que elija el usuario (por defecto, hoy).
"""
import datetime as dt
import json
import os
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.path.join(AQUI, "plantilla.html")
HOJA = "DetMaeSituAct"
CAMPOS_DIC = {  # columna excel -> clave del diccionario
    "dscdepa": "depa",
    "dsceje": "eje",
    "dscfnd": "fnd",
    "dscstaacso": "sta",
    "dscann70": "ann",
}


def fecha_int(v):
    """datetime -> YYYYMMDD (int). 0 si no hay fecha válida."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year * 10000 + v.month * 100 + v.day
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                d = dt.datetime.strptime(s, fmt)
                return d.year * 10000 + d.month * 100 + d.day
            except ValueError:
                pass
    return 0


def leer(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[-1]
    filas = ws.iter_rows(values_only=True)
    fuente = ""
    hdr = None
    datos = []
    for fila in filas:
        if hdr is None:
            celdas = [str(c).strip().lower() if c is not None else "" for c in fila]
            if "codsoc" in celdas and "fecnac" in celdas:
                hdr = celdas
                continue
            for c in fila:
                if isinstance(c, str) and c.lower().startswith("fuente"):
                    fuente = c.strip()
            continue
        if all(c is None for c in fila):
            continue
        datos.append(dict(zip(hdr, fila)))
    if hdr is None:
        raise SystemExit("No se encontró la cabecera (codsoc/fecnac) en la hoja %s" % ws.title)
    return fuente, datos


def compactar(datos):
    dic = {k: [] for k in CAMPOS_DIC.values()}
    idx = {k: {} for k in CAMPOS_DIC.values()}

    def cod(clave, valor):
        valor = (str(valor).strip() if valor is not None else "") or "(sin dato)"
        if valor not in idx[clave]:
            idx[clave][valor] = len(dic[clave])
            dic[clave].append(valor)
        return idx[clave][valor]

    filas = []
    for d in datos:
        edad = d.get("edad")
        try:
            edad = int(edad) if edad is not None and str(edad).strip() != "" else None
        except (TypeError, ValueError):
            edad = None
        filas.append([
            str(d.get("codsoc") or "").strip(),
            str(d.get("nombre") or "").strip(),
            str(d.get("ele") or "").strip(),
            str(d.get("s") or "").strip().upper(),
            fecha_int(d.get("fecnac")),
            edad,
        ] + [cod(clave, d.get(col)) for col, clave in CAMPOS_DIC.items()])
    return {
        "cols": ["cod", "nom", "dni", "s", "nac", "edadRep"] + list(CAMPOS_DIC.values()),
        "dic": dic,
        "filas": filas,
    }


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    ruta = sys.argv[1]
    salida = sys.argv[2] if len(sys.argv) > 2 else os.path.join(AQUI, "jubilaciones-amsp.html")
    fuente, datos = leer(ruta)
    paquete = compactar(datos)
    paquete["fuente"] = fuente
    paquete["archivo"] = os.path.basename(ruta)
    paquete["generado"] = dt.date.today().isoformat()
    paquete["entidad"] = "Asociación Mutualista Sanitaria del Perú"

    js = json.dumps(paquete, ensure_ascii=False, separators=(",", ":"))
    js = js.replace("</", "<\\/")  # nunca cerrar el <script> por accidente
    with open(PLANTILLA, encoding="utf-8") as f:
        html = f.read()
    marca = "/*__DATOS__*/"
    if marca not in html:
        raise SystemExit("La plantilla no tiene el marcador %s" % marca)
    html = html.replace(marca, "window.DATOS_INICIALES = " + js + ";", 1)
    with open(salida, "w", encoding="utf-8") as f:
        f.write(html)

    # Resumen rápido para verificar en consola.
    hoy = dt.date.today()
    f_ = sum(1 for r in paquete["filas"] if r[3] == "F")
    m_ = sum(1 for r in paquete["filas"] if r[3] == "M")
    sin_fecha = sum(1 for r in paquete["filas"] if not r[4])
    print("Fuente      :", fuente)
    print("Registros   :", len(paquete["filas"]), "(F %d / M %d, sin FECNAC %d)" % (f_, m_, sin_fecha))
    print("Salida      :", salida, "(%.1f MB)" % (os.path.getsize(salida) / 1e6))
    print("Generado el :", hoy.isoformat())


if __name__ == "__main__":
    main()
